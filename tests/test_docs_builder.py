"""Tests for agent/docs/builder.py style helpers."""

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

from agent.docs.builder import set_cn_font, set_cell_shading, add_page_break_before


def test_set_cn_font_sets_east_asia_attribute():
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("测试")
    set_cn_font(run, name="宋体", size_pt=12, bold=False)

    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    assert rFonts is not None
    assert rFonts.get(qn("w:eastAsia")) == "宋体"
    assert rFonts.get(qn("w:ascii")) == "宋体"
    assert run.font.size == Pt(12)
    assert run.font.bold is False


def test_set_cn_font_bold_flag():
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("标题")
    set_cn_font(run, name="黑体", size_pt=16, bold=True)
    assert run.font.bold is True


def test_set_cell_shading_adds_w_shd_element():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "E7E6E6")

    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    assert shd is not None
    assert shd.get(qn("w:fill")) == "E7E6E6"


def test_add_page_break_before_sets_property():
    doc = Document()
    p = doc.add_paragraph("Heading")
    add_page_break_before(p)

    pPr = p._p.get_or_add_pPr()
    pb = pPr.find(qn("w:pageBreakBefore"))
    assert pb is not None
    assert pb.get(qn("w:val")) == "true"


# ----- Task 5 tests (cover, header, revision table) -----

from datetime import date

from agent.docs.builder import build_cover, setup_page_and_header, build_revision_table


def _project_config_for_cover():
    return {
        "project": {
            "name": "图书管理系统",
            "version": "V1.0",
            "author": "测试作者",
            "code": "abc12",
        },
        "modules": [],
        "fields": {},
    }


def test_build_cover_adds_three_paragraphs():
    doc = Document()
    build_cover(doc, _project_config_for_cover())
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "图书管理系统" in texts
    assert "V1.0" in texts
    assert "使用手册" in texts


def test_setup_page_and_header_sets_header_text():
    doc = Document()
    setup_page_and_header(doc, name="图书管理系统", version="V1.0")
    section = doc.sections[0]
    header_text = section.header.paragraphs[0].text
    assert header_text == "图书管理系统V1.0使用手册"


def test_setup_page_and_header_sets_one_inch_margins():
    doc = Document()
    setup_page_and_header(doc, name="X", version="V1.0")
    section = doc.sections[0]
    from docx.shared import Inches
    assert section.top_margin == Inches(1)
    assert section.left_margin == Inches(1)
    assert section.right_margin == Inches(1)
    assert section.bottom_margin == Inches(1)


def test_build_revision_table_has_header_and_initial_row():
    doc = Document()
    today = date(2026, 6, 23)
    build_revision_table(doc, author="测试作者", today=today)
    tables = doc.tables
    assert len(tables) == 1
    tbl = tables[0]
    assert len(tbl.columns) == 4
    headers = [tbl.rows[0].cells[i].text for i in range(4)]
    assert headers == ["版本号", "修改日期", "作者", "修订内容"]
    first_row = [tbl.rows[1].cells[i].text for i in range(4)]
    assert first_row[0] == "V1.0"
    assert first_row[1] == "2026-06-23"
    assert first_row[2] == "测试作者"
    assert first_row[3] == "初始版本"


from agent.docs.schema import ProjectOutline, ModuleOutline
from agent.docs.builder import build_system_overview


def _outline_two_modules():
    return ProjectOutline(
        system_overview="图书管理系统包括图书管理和用户管理两个核心模块。",
        modules=[
            ModuleOutline(name="BookManagement", entity_name="图书",
                          overview_sentence="用于管理图书的入库、查询与下架。"),
            ModuleOutline(name="UserManagement", entity_name="用户",
                          overview_sentence="用于管理后台用户账号。"),
        ],
    )


def _config_two_modules():
    return {
        "project": {"name": "图书管理系统", "version": "V1.0", "author": "T"},
        "modules": [
            {"name": "BookManagement", "label": "图书管理", "description": "..."},
            {"name": "UserManagement", "label": "用户管理", "description": "..."},
        ],
        "fields": {},
    }


def test_build_system_overview_adds_h1_and_paragraphs():
    doc = Document()
    build_system_overview(doc, _config_two_modules(), _outline_two_modules())

    paragraphs = doc.paragraphs
    assert any("1. 系统概述" == p.text for p in paragraphs)
    assert any("系统包括图书管理和用户管理两个核心模块" in p.text for p in paragraphs)
    assert any(p.text.startswith("系统采用主流B/S") for p in paragraphs)


def test_build_system_overview_adds_core_features_per_module():
    doc = Document()
    build_system_overview(doc, _config_two_modules(), _outline_two_modules())

    texts = [p.text for p in doc.paragraphs]
    # 1.1 heading
    assert "1.1 核心功能" in texts
    # Each module should produce 4 capability lines
    assert any("新增图书信息" in t for t in texts)
    assert any("编辑图书信息" in t for t in texts)
    assert any("删除图书信息" in t for t in texts)
    assert any("查询图书信息" in t for t in texts)
    assert any("新增用户信息" in t for t in texts)


def test_build_system_overview_fallback_when_outline_missing():
    """If system_overview is empty, fallback to template."""
    doc = Document()
    cfg = _config_two_modules()
    outline = ProjectOutline(
        system_overview="",
        modules=_outline_two_modules().modules,
    )
    build_system_overview(doc, cfg, outline)
    texts = [p.text for p in doc.paragraphs]
    # Fallback should still mention project name
    assert any("图书管理系统包括" in t for t in texts)


from agent.docs.builder import build_compile_section, build_install_section


def test_build_compile_section_has_h1_and_intro():
    doc = Document()
    build_compile_section(doc)
    texts = [p.text for p in doc.paragraphs]
    assert "2. 软件编译" in texts
    assert any("Visual Studio Code" in t for t in texts)
    assert "软件开发界面：" in texts
    assert "编译成功画面：" in texts


def test_build_install_section_has_h1_and_four_subsections():
    doc = Document()
    build_install_section(doc)
    texts = [p.text for p in doc.paragraphs]
    assert "3. 软件安装" in texts
    assert "3.1 安装Python环境" in texts
    assert "3.2 安装系统依赖" in texts
    assert "3.3 启动系统" in texts
    assert "3.4 启动验证" in texts


def test_build_install_section_includes_install_steps():
    doc = Document()
    build_install_section(doc)
    texts = [p.text for p in doc.paragraphs]
    # Some step text should be present
    assert "python -m venv venv" in texts
    assert "pip install -r requirements.txt" in texts
    assert "python app.py" in texts


from agent.docs.builder import build_modules_section


def test_build_modules_section_single_module():
    doc = Document()
    cfg = _config_two_modules()
    outline = _outline_two_modules()
    # Test with only first module for predictability
    cfg["modules"] = [cfg["modules"][0]]
    outline.modules = [outline.modules[0]]
    build_modules_section(doc, cfg, outline)

    texts = [p.text for p in doc.paragraphs]
    assert "4. 功能模块" in texts
    # 4.1 should use module label
    assert "4.1 图书管理" in texts
    # 4.1.1 / 4.1.2 / 4.1.3
    assert "4.1.1 创建图书" in texts
    assert "4.1.2 查询图书" in texts
    assert "4.1.3 删除图书" in texts
    # Module intro and operation narratives
    assert any("点击【图书管理】" in t for t in texts)
    assert any("填写相关基础信息字段后点击【保存】" in t for t in texts)


def test_build_modules_section_two_modules_numbering():
    doc = Document()
    cfg = _config_two_modules()
    outline = _outline_two_modules()
    build_modules_section(doc, cfg, outline)

    texts = [p.text for p in doc.paragraphs]
    assert "4.1 图书管理" in texts
    assert "4.2 用户管理" in texts
    assert "4.2.1 创建用户" in texts
    assert "4.2.3 删除用户" in texts


def test_build_modules_section_uses_label_as_entity_fallback():
    """If outline lacks a module, fall back to label for entity."""
    doc = Document()
    cfg = _config_two_modules()
    # Empty outline.modules — should fall back
    outline = ProjectOutline(system_overview="x", modules=[])
    build_modules_section(doc, cfg, outline)

    texts = [p.text for p in doc.paragraphs]
    # Fallback uses label "图书管理" as entity in operation titles
    assert "4.1.1 创建图书管理" in texts


def test_build_modules_section_empty_modules():
    """No modules: only the chapter heading, no 4.x."""
    doc = Document()
    cfg = {"project": {"name": "X"}, "modules": [], "fields": {}}
    outline = ProjectOutline(system_overview="x", modules=[])
    build_modules_section(doc, cfg, outline)
    texts = [p.text for p in doc.paragraphs]
    assert "4. 功能模块" in texts
    # No 4.x subsections
    assert not any(t.startswith("4.1") for t in texts)


import tempfile
from pathlib import Path

from agent.docs.builder import build_tech_specs, build_appendix, build_manual


def test_build_tech_specs_has_hardware_and_software():
    doc = Document()
    build_tech_specs(doc)
    texts = [p.text for p in doc.paragraphs]
    assert "5. 技术规格" in texts
    assert "5.1 硬件要求" in texts
    assert "5.2 软件要求" in texts
    assert any("CPU" in t for t in texts)
    assert any("Python" in t for t in texts)


def test_build_appendix_has_glossary_and_references():
    doc = Document()
    build_appendix(doc)
    texts = [p.text for p in doc.paragraphs]
    assert "6. 附录" in texts
    assert "6.1 术语解释" in texts
    assert "6.2 参考资料" in texts
    assert any("SQLite3" in t for t in texts)
    assert any("Flask-Admin" in t for t in texts)


def test_build_manual_produces_docx_file():
    cfg = _config_two_modules()
    outline = _outline_two_modules()
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "test.docx"
        build_manual(cfg, outline, str(out))
        assert out.exists()
        # Verify by re-opening
        from docx import Document as OpenDoc
        doc = OpenDoc(str(out))
        texts = [p.text for p in doc.paragraphs]
        # Should contain all chapter headings in order
        assert "1. 系统概述" in texts
        assert "2. 软件编译" in texts
        assert "3. 软件安装" in texts
        assert "4. 功能模块" in texts
        assert "5. 技术规格" in texts
        assert "6. 附录" in texts


def test_build_manual_header_set_correctly():
    cfg = _config_two_modules()
    outline = _outline_two_modules()
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "test.docx"
        build_manual(cfg, outline, str(out))
        from docx import Document as OpenDoc
        doc = OpenDoc(str(out))
        header_text = doc.sections[0].header.paragraphs[0].text
        assert header_text == "图书管理系统V1.0使用手册"


def test_build_manual_chapter_order():
    cfg = _config_two_modules()
    outline = _outline_two_modules()
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "test.docx"
        build_manual(cfg, outline, str(out))
        from docx import Document as OpenDoc
        doc = OpenDoc(str(out))
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        # Verify chapter ordering by index
        idx_1 = next(i for i, t in enumerate(texts) if t == "1. 系统概述")
        idx_2 = next(i for i, t in enumerate(texts) if t == "2. 软件编译")
        idx_3 = next(i for i, t in enumerate(texts) if t == "3. 软件安装")
        idx_4 = next(i for i, t in enumerate(texts) if t == "4. 功能模块")
        idx_5 = next(i for i, t in enumerate(texts) if t == "5. 技术规格")
        idx_6 = next(i for i, t in enumerate(texts) if t == "6. 附录")
        assert idx_1 < idx_2 < idx_3 < idx_4 < idx_5 < idx_6
