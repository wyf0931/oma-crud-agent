"""python-docx assembly for user manual docx generation.

Style constants and helpers align with the pandora user-guide.docx
template.
"""

from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ===== Style constants (pt unless noted) =====
FONT_BODY = "宋体"
FONT_HEADING = "黑体"

SIZE_BODY = 12          # pt
SIZE_TITLE = 22         # pt (cover name + "使用手册")
SIZE_VERSION = 18       # pt (cover version)
SIZE_H1 = 16            # pt
SIZE_H2 = 14            # pt
SIZE_H3 = 12            # pt
SIZE_TABLE = 10.5       # pt
SIZE_HEADER = 9         # pt (page header)

TABLE_HEADER_FILL = "E7E6E6"


def set_cn_font(run, name: str = FONT_BODY, size_pt: float = SIZE_BODY, bold: bool = False) -> None:
    """Apply Chinese-aware font to a run.

    python-docx alone sets w:ascii/w:hAnsi but not w:eastAsia,
    so Word would fall back to Calibri for Chinese chars.
    """
    from docx.shared import Pt
    run.font.name = name
    run.font.size = None if size_pt is None else Pt(size_pt)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def set_cell_shading(cell, fill_hex: str) -> None:
    """Set table cell background color (hex without #)."""
    tcPr = cell._tc.get_or_add_tcPr()
    # Remove existing shading if any
    for existing in tcPr.findall(qn("w:shd")):
        tcPr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def add_page_break_before(paragraph) -> None:
    """Force this paragraph to start on a new page."""
    pPr = paragraph._p.get_or_add_pPr()
    # Remove existing pageBreakBefore if any
    for existing in pPr.findall(qn("w:pageBreakBefore")):
        pPr.remove(existing)
    pb = OxmlElement("w:pageBreakBefore")
    pb.set(qn("w:val"), "true")
    pPr.append(pb)


# ===== Page setup, header, cover, revision history =====

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from agent.docs import templates


def setup_page_and_header(doc, name: str, version: str) -> None:
    """Set 1-inch margins and add the standard page header."""
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.text = templates.HEADER_TEXT.format(name=name, version=version)
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in hp.runs:
        set_cn_font(run, name=FONT_BODY, size_pt=SIZE_HEADER, bold=False)


def build_cover(doc, project_config: dict) -> None:
    """Cover page: project name -> version -> '使用手册'."""
    project = project_config["project"]
    name = project["name"]
    version = project["version"]

    # Project name (large title)
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1 = p1.add_run(name)
    set_cn_font(run1, name=FONT_HEADING, size_pt=SIZE_TITLE, bold=True)

    # Version
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(version)
    set_cn_font(run2, name=FONT_HEADING, size_pt=SIZE_VERSION, bold=False)

    # "使用手册"
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run3 = p3.add_run("使用手册")
    set_cn_font(run3, name=FONT_HEADING, size_pt=SIZE_TITLE, bold=True)


def build_revision_table(doc, author: str, today) -> None:
    """文件修订记录 table: 4 cols, header + one initial row."""
    from datetime import date as _date
    if isinstance(today, _date):
        date_str = today.strftime("%Y-%m-%d")
    else:
        date_str = str(today)

    headers = ["版本号", "修改日期", "作者", "修订内容"]
    table = doc.add_table(rows=2, cols=4)
    table.style = "Table Grid"

    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_cn_font(run, name=FONT_BODY, size_pt=SIZE_TABLE, bold=True)
        set_cell_shading(cell, TABLE_HEADER_FILL)

    # Initial revision row
    initial = ["V1.0", date_str, author, templates.REVISION_INITIAL_CONTENT]
    for i, v in enumerate(initial):
        cell = table.rows[1].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(v)
        set_cn_font(run, name=FONT_BODY, size_pt=SIZE_TABLE, bold=False)


# ===== Chapter 1: System Overview =====

def _add_heading(doc, text: str, level: int, page_break: bool = False) -> None:
    """Add a heading with Chinese font and optional page break."""
    p = doc.add_paragraph()
    if page_break:
        add_page_break_before(p)
    run = p.add_run(text)
    if level == 1:
        set_cn_font(run, name=FONT_HEADING, size_pt=SIZE_H1, bold=True)
    elif level == 2:
        set_cn_font(run, name=FONT_HEADING, size_pt=SIZE_H2, bold=True)
    else:
        set_cn_font(run, name=FONT_HEADING, size_pt=SIZE_H3, bold=True)


def _add_body(doc, text: str) -> None:
    """Add a body paragraph with 宋体 12pt."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_cn_font(run, name=FONT_BODY, size_pt=SIZE_BODY, bold=False)


def build_system_overview(doc, project_config: dict, outline) -> None:
    """Chapter 1: 系统概述 + 1.1 核心功能."""
    project = project_config["project"]
    modules = project_config.get("modules", [])
    name = project["name"]

    _add_heading(doc, "1. 系统概述", level=1, page_break=True)

    # Paragraph 1: system overview (LLM-generated or fallback)
    overview = outline.system_overview.strip() if outline and outline.system_overview else ""
    if not overview:
        module_list_str = "、".join(m["label"] for m in modules)
        overview = templates.SYSTEM_OVERVIEW_FALLBACK.format(
            name=name, module_list=module_list_str
        )
    _add_body(doc, overview)

    # Paragraph 2: fixed tech stack description
    _add_body(doc, templates.TECH_STACK_PARAGRAPH)

    # 1.1 核心功能
    _add_heading(doc, "1.1 核心功能", level=2)

    # Map module.name -> entity_name from outline
    entity_map = {}
    if outline and outline.modules:
        for m in outline.modules:
            entity_map[m.name] = m.entity_name

    for m in modules:
        # Module label as sub-heading (plain paragraph, bold)
        p_label = doc.add_paragraph()
        run_label = p_label.add_run(m["label"])
        set_cn_font(run_label, name=FONT_BODY, size_pt=SIZE_BODY, bold=True)

        entity = entity_map.get(m["name"], m["label"])
        for tpl in templates.CORE_FEATURE_TEMPLATES:
            _add_body(doc, tpl.format(entity=entity))


# ===== Chapter 2: 软件编译 (fixed) =====

def build_compile_section(doc) -> None:
    _add_heading(doc, "2. 软件编译", level=1, page_break=True)
    _add_body(doc, templates.SECTION_COMPILE_INTRO)
    _add_body(doc, templates.SECTION_COMPILE_DEV_UI)
    _add_body(doc, "")  # placeholder gap for screenshot
    _add_body(doc, templates.SECTION_COMPILE_BUILD_OK)
    _add_body(doc, "")  # placeholder gap


# ===== Chapter 3: 软件安装 (fixed) =====

def build_install_section(doc) -> None:
    _add_heading(doc, "3. 软件安装", level=1, page_break=True)

    # 3.1
    _add_heading(doc, "3.1 安装Python环境", level=2)
    _add_body(doc, templates.INSTALL_PYTHON)

    # 3.2
    _add_heading(doc, "3.2 安装系统依赖", level=2)
    for step in templates.INSTALL_DEPS_STEPS:
        _add_body(doc, step)

    # 3.3
    _add_heading(doc, "3.3 启动系统", level=2)
    _add_body(doc, templates.INSTALL_START_INTRO)
    _add_body(doc, templates.INSTALL_START_CMD)
    _add_body(doc, templates.INSTALL_START_OK)
    _add_body(doc, "")  # screenshot gap

    # 3.4
    _add_heading(doc, "3.4 启动验证", level=2)
    _add_body(doc, templates.INSTALL_VERIFY_INTRO)
    _add_body(doc, "")  # screenshot gap
    _add_body(doc, templates.INSTALL_VERIFY_OK)
    _add_body(doc, "")  # screenshot gap


# ===== Chapter 4: 功能模块 (dynamic per module) =====

def build_modules_section(doc, project_config: dict, outline) -> None:
    """Chapter 4: one 4.x section per module, each with 创建/查询/删除 subsections."""
    modules = project_config.get("modules", [])

    _add_heading(doc, "4. 功能模块", level=1, page_break=True)

    # Map module.name -> entity_name
    entity_map = {}
    if outline and outline.modules:
        for m in outline.modules:
            entity_map[m.name] = m.entity_name

    for idx, module in enumerate(modules, start=1):
        label = module["label"]
        # Fallback: use label as entity if outline missing
        entity = entity_map.get(module["name"], label)

        # 4.x heading
        _add_heading(doc, f"4.{idx} {label}", level=2)

        # Module intro
        _add_body(doc, templates.MODULE_INTRO.format(label=label, entity=entity))
        _add_body(doc, "")  # screenshot gap

        # 4.x.1 创建
        _add_heading(doc, f"4.{idx}.1 创建{entity}", level=3)
        _add_body(doc, templates.CREATE_STEP1.format(label=label, entity=entity))
        _add_body(doc, "")  # screenshot
        _add_body(doc, templates.CREATE_STEP2)
        _add_body(doc, "")  # screenshot
        _add_body(doc, templates.CREATE_STEP3.format(entity=entity))
        _add_body(doc, "")  # screenshot

        # 4.x.2 查询
        _add_heading(doc, f"4.{idx}.2 查询{entity}", level=3)
        _add_body(doc, templates.QUERY_STEP1.format(label=label))
        _add_body(doc, "")  # screenshot
        _add_body(doc, templates.QUERY_STEP2)
        _add_body(doc, "")  # screenshot

        # 4.x.3 删除
        _add_heading(doc, f"4.{idx}.3 删除{entity}", level=3)
        _add_body(doc, templates.DELETE_STEP1.format(label=label))
        _add_body(doc, templates.DELETE_STEP2)
        _add_body(doc, "")  # screenshot
        _add_body(doc, templates.DELETE_STEP3.format(entity=entity))


# ===== Chapter 5: 技术规格 (fixed) =====

def build_tech_specs(doc) -> None:
    _add_heading(doc, "5. 技术规格", level=1, page_break=True)

    _add_heading(doc, "5.1 硬件要求", level=2)
    for req in templates.TECH_HARDWARE_REQS:
        _add_body(doc, req)

    _add_heading(doc, "5.2 软件要求", level=2)
    for req in templates.TECH_SOFTWARE_REQS:
        _add_body(doc, req)


# ===== Chapter 6: 附录 (fixed) =====

def build_appendix(doc) -> None:
    _add_heading(doc, "6. 附录", level=1, page_break=True)

    _add_heading(doc, "6.1 术语解释", level=2)
    for term, explanation in templates.GLOSSARY:
        _add_body(doc, f"{term}：{explanation}")

    _add_heading(doc, "6.2 参考资料", level=2)
    for name, url in templates.REFERENCES:
        _add_body(doc, f"{name} Documentation: {url}")


# ===== Orchestrator =====

def build_manual(project_config: dict, outline, output_path: str) -> str:
    """Assemble the full user manual docx and save to output_path.

    Args:
        project_config: project_config dict (matches session state.project_config)
        outline: ProjectOutline instance (may be empty/fallback)
        output_path: absolute path to save .docx

    Returns:
        The output_path (for convenience).
    """
    from docx import Document
    from datetime import date
    from pathlib import Path

    project = project_config["project"]
    name = project["name"]
    version = project["version"]
    author = project.get("author", "")

    doc = Document()

    # Page setup + header
    setup_page_and_header(doc, name=name, version=version)

    # Cover
    build_cover(doc, project_config)

    # Revision history table
    build_revision_table(doc, author=author, today=date.today())

    # Chapter 1: 系统概述 + 1.1 核心功能
    build_system_overview(doc, project_config, outline)

    # Chapter 2: 软件编译
    build_compile_section(doc)

    # Chapter 3: 软件安装
    build_install_section(doc)

    # Chapter 4: 功能模块 (dynamic)
    build_modules_section(doc, project_config, outline)

    # Chapter 5: 技术规格
    build_tech_specs(doc)

    # Chapter 6: 附录
    build_appendix(doc)

    # Ensure parent dir exists
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path
